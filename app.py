import io
import os
import re
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pdfplumber
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, portrait
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
import streamlit as st

# ==============================================================================
# 1. KONSTANTA & DETEKSI BANK / BULAN
# ==============================================================================
MONTH_NAMES_ID = [
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember",
]

MONTH_MAP = {
    "JAN": "Januari",
    "JANUARY": "Januari",
    "FEB": "Februari",
    "FEBRUARY": "Februari",
    "MAR": "Maret",
    "MARCH": "Maret",
    "APR": "April",
    "APRIL": "April",
    "MAY": "Mei",
    "MEI": "Mei",
    "JUN": "Juni",
    "JUNE": "Juni",
    "JUL": "Juli",
    "JULY": "Juli",
    "AUG": "Agustus",
    "AUGUST": "Agustus",
    "AGU": "Agustus",
    "SEP": "September",
    "SEPTEMBER": "September",
    "OCT": "Oktober",
    "OCTOBER": "Oktober",
    "OKT": "Oktober",
    "NOV": "November",
    "NOVEMBER": "November",
    "DEC": "Desember",
    "DECEMBER": "Desember",
    "DES": "Desember",
}


def detect_bank_and_month(full_text, filename=""):
  header_scope = full_text[:1500].upper()
  full_upper = (full_text + " " + filename).upper()

  # 1. Deteksi Bank Berdasarkan Header Dokumen
  if any(k in header_scope for k in ["BANK CENTRAL ASIA", "REKENING GIRO"]) or (
      "BCA" in header_scope and "MUTASI" in header_scope
  ) or ("BCA" in filename.upper() and ("REKENING" in header_scope or "GIRO" in header_scope)):
    bank = "BCA"
  elif any(k in header_scope for k in ["BANK SYARIAH INDONESIA", "BSI", "BSINET"]) or ("BSI" in filename.upper()):
    bank = "BSI"
  elif any(
      k in header_scope
      for k in [
          "BANK RAKYAT INDONESIA",
          "SETULUS HATI",
          "LAPORAN TRANSAKSI FINANSIAL",
      ]
  ) or ("IBIZ" in header_scope and "BRI" in full_upper) or ("BRI" in filename.upper() and "TRANSAKSI" in header_scope):
    bank = "BRI"
  elif any(
      k in header_scope
      for k in [
          "BANK NEGARA INDONESIA",
          "BNI DIRECT",
          "LEDGER BALANCE",
      ]
  ) or ("BNI" in header_scope and "ACCOUNT STATEMENT" in header_scope) or ("BNI" in filename.upper()):
    bank = "BNI"
  elif any(k in header_scope for k in ["BANK BJB", "BANK JABAR", "BJB"]) or (
      "ACCOUNT ACTIVITY" in header_scope and "GIRO UMUM" in header_scope
  ) or ("BJB" in filename.upper()):
    bank = "BJB"
  elif any(k in header_scope for k in ["BANK MANDIRI", "KOPRA", "MCM"]) or (
      "MANDIRI" in header_scope and "REKENING" in header_scope
  ) or ("MANDIRI" in filename.upper()):
    bank = "MANDIRI"
  elif re.search(r"\b(?:BANK\s+PERMATA|PERMATABANK|PERMATA\s+BANK)\b", header_scope) or ("PERMATA" in filename.upper()):
    bank = "PERMATA"
  elif re.search(r"\b(?:BANK\s+NOBU|NOBU\s+BANK|NATIONALNOBU)\b", header_scope) or ("NOBU" in filename.upper()):
    bank = "NOBU"
  elif "BANK SYARIAH INDONESIA" in full_upper or "BSI" in filename.upper():
    bank = "BSI"
  elif "BANK CENTRAL ASIA" in full_upper or "BCA" in filename.upper():
    bank = "BCA"
  elif "BANK RAKYAT INDONESIA" in full_upper or "BRI" in filename.upper():
    bank = "BRI"
  elif "BANK NEGARA INDONESIA" in full_upper or "BNI" in filename.upper():
    bank = "BNI"
  elif "BANK MANDIRI" in full_upper or "MANDIRI" in filename.upper():
    bank = "MANDIRI"
  elif "BANK BJB" in full_upper or "BJB" in filename.upper():
    bank = "BJB"
  elif re.search(r"\b(?:BANK\s+PERMATA|PERMATABANK)\b", full_upper):
    bank = "PERMATA"
  elif re.search(r"\b(?:BANK\s+NOBU|NOBU\s+BANK)\b", full_upper):
    bank = "NOBU"
  else:
    bank = "UMUM"

  detected_month = "Bulan"

  # 2. Deteksi Bulan Dari Header Periode
  m_periode = re.search(r"PERIODE\s*:\s*([A-Za-z]+)\s+\d{4}", header_scope)
  if m_periode:
    m_cand = m_periode.group(1).upper()
    if m_cand in MONTH_MAP:
      detected_month = MONTH_MAP[m_cand]
    else:
      for m_id in MONTH_NAMES_ID:
        if m_id.upper() == m_cand:
          detected_month = m_id
          break

  if detected_month == "Bulan":
    m_range_str = re.search(
        r"(?:Period|Date|Periode)\s*[:\s]*\d{2}\s+([A-Za-z]{3,9})\s+\d{2,4}",
        header_scope,
        re.IGNORECASE,
    )
    if m_range_str:
      m_cand = m_range_str.group(1).upper()
      if m_cand in MONTH_MAP:
        detected_month = MONTH_MAP[m_cand]
      else:
        for m_id in MONTH_NAMES_ID:
          if m_id.upper() == m_cand:
            detected_month = m_id
            break

  if detected_month == "Bulan":
    m_bni = re.search(
        r"(?:Period|Periode)\s*[:\s]*\d{2}-([A-Za-z]{3})-\d{2,4}",
        header_scope,
        re.IGNORECASE,
    )
    if m_bni:
      m_cand = m_bni.group(1).upper()
      if m_cand in MONTH_MAP:
        detected_month = MONTH_MAP[m_cand]

  if detected_month == "Bulan":
    m_slash = re.search(
        r"(?:Periode\s+Transaksi|Period|Periode)\s*[:\s]*\d{2}/(\d{2})/\d{2,4}",
        header_scope,
        re.IGNORECASE,
    )
    if m_slash:
      m_num = int(m_slash.group(1))
      if 1 <= m_num <= 12:
        detected_month = MONTH_NAMES_ID[m_num - 1]

  if detected_month == "Bulan":
    fn_upper = filename.upper()
    for m_id in MONTH_NAMES_ID:
      if m_id.upper() in fn_upper:
        detected_month = m_id
        break

  return bank, detected_month


def extract_metadata_from_pdf(all_text, first_page_text, bank):
  meta = {
      "cabang": "",
      "nama_cust": "",
      "no_rekening": "",
      "nama_bank": bank,
      "nama_pemegang_rek": "",
  }
  p1 = first_page_text

  if bank == "BCA":
    cbg_m = re.search(r"REKENING GIRO\s*\n\s*([^\n]+)", p1)
    if cbg_m:
      meta["cabang"] = cbg_m.group(1).strip()
    cust_m = re.search(r"KCP[^\n]*\n+([^\n]+?)(?:\s+NO\.\s*REKENING|\n)", p1)
    if cust_m:
      meta["nama_cust"] = cust_m.group(1).strip()
      meta["nama_pemegang_rek"] = meta["nama_cust"]
    rek_m = re.search(r"NO\.\s*REKENING\s*:\s*(\d+)", p1)
    if rek_m:
      meta["no_rekening"] = rek_m.group(1).strip()

  elif bank == "BJB":
    cbg_m = re.search(r"Branch\s*:\s*(?:[0-9]+\s*-\s*)?([^\n]+)", p1)
    if cbg_m:
      meta["cabang"] = cbg_m.group(1).strip()
    cust_m = re.search(r"Corporate\s*:\s*([^\n]+)", p1)
    if cust_m:
      meta["nama_cust"] = cust_m.group(1).strip()
      meta["nama_pemegang_rek"] = meta["nama_cust"]
    rek_m = re.search(r"Account\s*:\s*(\d+)", p1)
    if rek_m:
      meta["no_rekening"] = rek_m.group(1).strip()

  elif bank == "BNI":
    rek_m = re.search(
        r"Account\s+No\.\s*:\s*(\d+)\s*/\s*([^\n]+?)(?:\s*PT\(IDR\)|\(IDR\)|\n)",
        p1,
    )
    if rek_m:
      meta["no_rekening"] = rek_m.group(1).strip()
      cust_name = rek_m.group(2).strip()
      if not cust_name.endswith("PT") and "PT" in p1:
        cust_name += " PT"
      meta["nama_cust"] = cust_name
      meta["nama_pemegang_rek"] = cust_name
    else:
      rek_m2 = re.search(r"Account\s+No\.\s*:\s*(\d+)", p1)
      if rek_m2:
        meta["no_rekening"] = rek_m2.group(1).strip()

  elif bank == "BRI":
    cbg_m = re.search(r"Unit Kerja[^\n:]*:\s*([^\n]+)", p1)
    if cbg_m:
      meta["cabang"] = cbg_m.group(1).strip()
    cust_m = re.search(
        r"Statement Date\s*\n\s*([^\n]+?)\s*(?:Periode Transaksi|\n)", p1
    )
    if not cust_m:
      cust_m = re.search(r"Kepada Yth\.[^\n]*\n+([^\n]+)", p1)
    if cust_m:
      meta["nama_cust"] = cust_m.group(1).strip()
      meta["nama_pemegang_rek"] = meta["nama_cust"]
    rek_m = re.search(r"No\.\s*Rekening[^\n:]*:\s*(\d+)", p1)
    if rek_m:
      meta["no_rekening"] = rek_m.group(1).strip()

  elif bank == "BSI":
    rek_m = re.search(r"Account\s*:\s*(\d+)", p1)
    if rek_m:
      meta["no_rekening"] = rek_m.group(1).strip()
    cust_m = re.search(
        r"Account Statement\s*-\s*\d+\s*-\s*([^\n\d]+?)(?:\s+\d{2}\s+[A-Za-z]+|\n)",
        p1,
    )
    if not cust_m:
      cust_m = re.search(r"^([^\n]+?)\s*/\s*IDR", p1, re.MULTILINE)
    if cust_m:
      meta["nama_cust"] = cust_m.group(1).strip()
      meta["nama_pemegang_rek"] = meta["nama_cust"]
    cbg_m = re.search(r"Branch\s*:\s*([^\n]+)", p1)
    if cbg_m:
      meta["cabang"] = cbg_m.group(1).strip()

  elif bank == "MANDIRI":
    rek_m = re.search(
        r"(?:No\.\s*Rekening|Account Number)\s*:\s*(\d+)", p1
    )
    if not rek_m:
      rek_m = re.search(r"Account No\s+(\d+)", p1)
    if rek_m:
      meta["no_rekening"] = rek_m.group(1).strip()
    cust_m = re.search(
        r"(?:Nama Nasabah|Customer Name)\s*:\s*([^\n]+)", p1
    )
    if not cust_m:
      cust_m = re.search(
          r"Account No\s+\d+\s+IDR\s+([^\n]+?)\s+[A-Z0-9\s\-]+-", p1
      )
    if cust_m:
      meta["nama_cust"] = cust_m.group(1).strip()
      meta["nama_pemegang_rek"] = meta["nama_cust"]
    cbg_m = re.search(r"Branch\s+([^\n]+)", p1)
    if cbg_m:
      meta["cabang"] = cbg_m.group(1).strip()

  return meta


# ==============================================================================
# 2. PARSER SPESIFIK BANK
# ==============================================================================
def parse_bni_clean(pdf, all_text):
  ledger_m = re.search(
      r"Ledger\s+Balance\s*:\s*([\d,]+\.\d{2})", all_text, re.IGNORECASE
  )
  opening_balance = (
      float(ledger_m.group(1).replace(",", "")) if ledger_m else None
  )

  raw_items = []
  for page in pdf.pages:
    words = page.extract_words(x_tolerance=3, y_tolerance=3)
    if not words:
      continue
    p_width = float(page.width)
    x_bal_min = p_width * 0.84

    lines_dict = {}
    for w in words:
      y_mid = round(w["top"] / 4.0) * 4.0
      lines_dict.setdefault(y_mid, []).append(w)

    for y in sorted(lines_dict.keys()):
      lw = sorted(lines_dict[y], key=lambda x: x["x0"])
      line_str = " ".join(w["text"] for w in lw).strip()

      if any(
          k in line_str.upper()
          for k in [
              "ACCOUNT STATEMENT",
              "POSTING DATE",
              "TOTAL DEBET",
              "TOTAL CREDIT",
              "PAGE :",
          ]
      ):
        continue

      tgl_m = re.search(r"\b(\d{2}/\d{2})/\d{4}\b", line_str)
      tgl_str = tgl_m.group(1) if tgl_m else ""

      bal_words = [
          w
          for w in lw
          if w["x0"] >= x_bal_min - 15 and re.match(r"^[\d,]+\.\d{2}$", w["text"])
      ]
      if not bal_words:
        continue

      sal_val = float(bal_words[-1]["text"].replace(",", ""))
      if "LEDGER BALANCE" not in line_str.upper():
        raw_items.append(
            {"date": tgl_str, "saldo": sal_val, "page": page.page_number}
        )

  cur_d = ""
  for it in raw_items:
    if it["date"]:
      cur_d = it["date"]
    else:
      it["date"] = cur_d

  tx_records = []
  balances = []
  prev_sal = (
      opening_balance
      if opening_balance is not None
      else (raw_items[0]["saldo"] if raw_items else 0.0)
  )

  for it in raw_items:
    cur_sal = it["saldo"]
    delta = round(cur_sal - prev_sal, 2)
    if delta == 0:
      continue

    deb = abs(delta) if delta < 0 else 0.0
    krd = delta if delta > 0 else 0.0

    balances.append(cur_sal)
    tx_records.append(
        {"date": it["date"], "debet": deb, "kredit": krd, "saldo": cur_sal}
    )
    prev_sal = cur_sal

  return tx_records, balances, opening_balance


def parse_permata_clean(pages_text, all_text):
  op_m = re.search(
      r"Opening\s+Balance\s*[:\|]?\s*([\d,]+\.\d{2})", all_text, re.IGNORECASE
  )
  tot_deb_m = re.search(
      r"Total\s+Debit\s*[:\|]?\s*([\d,]+\.\d{2})", all_text, re.IGNORECASE
  )
  tot_krd_m = re.search(
      r"Total\s+Credit\s*[:\|]?\s*([\d,]+\.\d{2})", all_text, re.IGNORECASE
  )

  opening_balance = float(op_m.group(1).replace(",", "")) if op_m else 0.0
  mutasi_db_off = (
      float(tot_deb_m.group(1).replace(",", "")) if tot_deb_m else None
  )
  mutasi_cr_off = (
      float(tot_krd_m.group(1).replace(",", "")) if tot_krd_m else None
  )

  raw_tx = []
  for txt in pages_text:
    for line in txt.split("\n"):
      line_str = line.strip()
      if not line_str or any(
          k in line_str.upper()
          for k in [
              "TRANSACTION HISTORY",
              "TRANSACTION DATE",
              "VALUE DATE",
              "OPENING BALANCE",
              "CLOSING BALANCE",
          ]
      ):
        continue

      tgl_m = re.match(r"^(\d{2})\s+([A-Za-z]{3})\s+\d{4}\b", line_str)
      if not tgl_m:
        continue

      tgl_str = f"{tgl_m.group(1)}/{tgl_m.group(2).upper()}"
      amt_m = re.search(r"(-?\s*[\d,]+\.\d{2}-?)\s*$", line_str)
      if not amt_m:
        continue

      amt_str = amt_m.group(1).strip()
      is_minus = "-" in amt_str
      cleaned_num = re.sub(r"[^\d\.]", "", amt_str.replace(",", ""))

      try:
        num_val = float(cleaned_num)
      except ValueError:
        continue

      deb = num_val if is_minus else 0.0
      krd = num_val if not is_minus else 0.0
      raw_tx.append({"date": tgl_str, "debet": deb, "kredit": krd})

  raw_tx.reverse()
  tx_records = []
  balances = []
  running_s = opening_balance

  for r in raw_tx:
    running_s = round(running_s + r["kredit"] - r["debet"], 2)
    balances.append(running_s)
    tx_records.append({
        "date": r["date"],
        "debet": r["debet"],
        "kredit": r["kredit"],
        "saldo": running_s,
    })

  freq_db = sum(1 for r in tx_records if r["debet"] > 0)
  freq_cr = sum(1 for r in tx_records if r["kredit"] > 0)
  mutasi_db = (
      mutasi_db_off
      if mutasi_db_off is not None
      else sum(r["debet"] for r in tx_records)
  )
  mutasi_cr = (
      mutasi_cr_off
      if mutasi_cr_off is not None
      else sum(r["kredit"] for r in tx_records)
  )

  return (
      tx_records,
      balances,
      opening_balance,
      freq_db,
      freq_cr,
      mutasi_db,
      mutasi_cr,
  )


def parse_mandiri_clean(pages_text, all_text, pdf=None):
  # 1. Opening Balance Mandiri Kopra / Rekening Giro
  op_m = re.search(
      r"Opening\s+Balance[^\d]*([\d,]+\.\d{2})",
      all_text[:3000],
      re.IGNORECASE,
  )
  opening_balance = float(op_m.group(1).replace(",", "")) if op_m else None

  tx_records = []
  balances = []

  # Coba baca via extract_tables bila ada (Mandiri Kopra terstruktur rapi)
  table_read_success = False
  if pdf:
    for page in pdf.pages:
      tables = page.extract_tables()
      for table in tables:
        for row in table:
          if not row or len(row) < 5:
            continue
          row_str = " ".join([str(c) for c in row if c])
          tgl_m = re.search(r"\b(\d{2}\s+[A-Za-z]{3}\s+\d{4})", row_str)
          if tgl_m:
            try:
              # Normalisasi angka dari 3 kolom terakhir: Debit, Credit, Balance
              s_val = float(str(row[-1]).replace(",", "").strip())
              k_val = float(str(row[-2]).replace(",", "").strip())
              d_val = float(str(row[-3]).replace(",", "").strip())

              tx_records.append({
                  "date": tgl_m.group(1)[:6],
                  "debet": d_val,
                  "kredit": k_val,
                  "saldo": s_val,
              })
              balances.append(s_val)
              table_read_success = True
            except (ValueError, IndexError):
              continue

  # Fallback text parser jika bukan format tabel Kopra
  if not table_read_success:
    active_date = ""
    for txt in pages_text:
      for line in txt.split("\n"):
        line_str = line.strip()
        if not line_str or any(
            k in line_str.upper()
            for k in [
                "ACCOUNT STATEMENT REPORT",
                "CLOSING BALANCE",
                "TOTAL DEBIT",
            ]
        ):
          continue

        tgl_match = re.search(
            r"\b(\d{2}/\d{2}/\d{4}|\d{2}\s+[A-Za-z]{3}\s+\d{4})\b", line_str
        )
        if tgl_match:
          active_date = tgl_match.group(1)[:5]

        nums_m = re.findall(r"([\d,]+\.\d{2})", line_str)
        if len(nums_m) >= 3:
          last_3 = [float(n.replace(",", "")) for n in nums_m[-3:]]
          d_val, k_val, s_val = last_3[0], last_3[1], last_3[2]
          if (
              (d_val > 0 or k_val > 0)
              and s_val > 0
              and "OPENING BALANCE" not in line_str.upper()
          ):
            balances.append(s_val)
            tx_records.append({
                "date": active_date,
                "debet": d_val,
                "kredit": k_val,
                "saldo": s_val,
            })

  if opening_balance is None and tx_records:
    opening_balance = round(
        tx_records[0]["saldo"]
        + tx_records[0]["debet"]
        - tx_records[0]["kredit"],
        2,
    )

  freq_db = sum(1 for r in tx_records if r["debet"] > 0)
  freq_cr = sum(1 for r in tx_records if r["kredit"] > 0)
  mutasi_db = sum(r["debet"] for r in tx_records)
  mutasi_cr = sum(r["kredit"] for r in tx_records)

  return (
      tx_records,
      balances,
      opening_balance,
      freq_db,
      freq_cr,
      mutasi_db,
      mutasi_cr,
  )


def parse_bsi_clean(pdf, all_text):
  op_m = re.search(
      r"Opening\s+Balance[\s\S]{0,30}?(?:IDR)?\s*([\d,]+\.\d{2})",
      all_text,
      re.IGNORECASE,
  )
  opening_balance = float(op_m.group(1).replace(",", "")) if op_m else 0.0

  tot_deb_m = re.search(
      r"Total\s+Debit\s+Amount\s*[:\|]?\s*(?:IDR)?\s*([\d,]+\.\d{2})",
      all_text,
      re.IGNORECASE,
  )
  tot_krd_m = re.search(
      r"Total\s+Credit\s+Amount\s*[:\|]?\s*(?:IDR)?\s*([\d,]+\.\d{2})",
      all_text,
      re.IGNORECASE,
  )
  tot_deb_rec_m = re.search(
      r"Total\s+Debit\s+Record\s*[:\|]?\s*(\d+)", all_text, re.IGNORECASE
  )
  tot_crd_rec_m = re.search(
      r"Total\s+Credit\s+Record\s*[:\|]?\s*(\d+)", all_text, re.IGNORECASE
  )

  mutasi_db_off = (
      float(tot_deb_m.group(1).replace(",", "")) if tot_deb_m else None
  )
  mutasi_cr_off = (
      float(tot_krd_m.group(1).replace(",", "")) if tot_krd_m else None
  )
  freq_db_off = int(tot_deb_rec_m.group(1)) if tot_deb_rec_m else None
  freq_cr_off = int(tot_crd_rec_m.group(1)) if tot_crd_rec_m else None

  tx_records = []
  balances = []
  active_date = ""

  for page in pdf.pages:
    txt = page.extract_text() or ""
    for line in txt.split("\n"):
      line_str = line.strip()
      if not line_str or any(
          k in line_str.upper()
          for k in [
              "ACCOUNT STATEMENT",
              "OPENING BALANCE",
              "TOTAL DEBIT",
              "TOTAL CREDIT",
              "CLOSING BALANCE",
          ]
      ):
        continue

      tgl_m = re.search(r"\b(\d{4})-(\d{2})-(\d{2})\b", line_str)
      if tgl_m:
        active_date = f"{tgl_m.group(3)}/{tgl_m.group(2)}"

      line_match = re.search(
          r"([\d,]+\.\d{2})\s+(DB|CR)\s+([\d,]+\.\d{2})", line_str
      )
      if line_match:
        val_amt = float(line_match.group(1).replace(",", ""))
        flag = line_match.group(2)
        val_bal = float(line_match.group(3).replace(",", ""))
        deb = val_amt if flag == "DB" else 0.0
        krd = val_amt if flag == "CR" else 0.0

        if val_bal > 0 and (deb > 0 or krd > 0):
          balances.append(val_bal)
          tx_records.append({
              "date": active_date,
              "debet": deb,
              "kredit": krd,
              "saldo": val_bal,
          })

  freq_db = (
      freq_db_off
      if freq_db_off is not None
      else sum(1 for r in tx_records if r["debet"] > 0)
  )
  freq_cr = (
      freq_cr_off
      if freq_cr_off is not None
      else sum(1 for r in tx_records if r["kredit"] > 0)
  )
  mutasi_db = (
      mutasi_db_off
      if mutasi_db_off is not None
      else sum(r["debet"] for r in tx_records)
  )
  mutasi_cr = (
      mutasi_cr_off
      if mutasi_cr_off is not None
      else sum(r["kredit"] for r in tx_records)
  )

  return (
      tx_records,
      balances,
      opening_balance,
      freq_db,
      freq_cr,
      mutasi_db,
      mutasi_cr,
  )


def parse_bca_clean(pdf, all_text):
  sa_m = re.search(r"SALDO AWAL\s*:\s*([\d,\.]+)", all_text)
  cr_m = re.search(r"MUTASI CR\s*:\s*([\d,\.]+)\s+(\d+)", all_text)
  db_m = re.search(r"MUTASI DB\s*:\s*([\d,\.]+)\s+(\d+)", all_text)

  opening_bal = float(sa_m.group(1).replace(",", "")) if sa_m else None
  mutasi_cr_off = float(cr_m.group(1).replace(",", "")) if cr_m else None
  freq_cr_off = int(cr_m.group(2)) if cr_m else None
  mutasi_db_off = float(db_m.group(1).replace(",", "")) if db_m else None
  freq_db_off = int(db_m.group(2)) if db_m else None

  tx_records = []
  active_date = ""

  for page in pdf.pages:
    txt = page.extract_text() or ""
    lines = txt.split("\n")
    for line in lines:
      line_str = line.strip()
      if not line_str or any(
          k in line_str.upper()
          for k in [
              "REKENING GIRO",
              "NO. REKENING",
              "HALAMAN :",
              "PERIODE :",
              "MATA UANG",
              "CATATAN:",
              "TANGGAL KETERANGAN",
              "MUTASI CR",
              "MUTASI DB",
              "SALDO AKHIR",
              "BERSAMBUNG",
          ]
      ):
        continue

      if "SALDO AWAL" in line_str.upper():
        if opening_bal is None:
          nums = re.findall(r"([\d,]+\.\d{2})", line_str)
          if nums:
            opening_bal = float(nums[-1].replace(",", ""))
        continue

      m_date = re.match(r"^(\d{2}/\d{2})\s+", line_str)
      if m_date:
        active_date = m_date.group(1)

      if not active_date:
        continue

      # Baris Debet (dengan akhiran DB)
      m_db = re.search(
          r"([\d,]+\.\d{2})\s+DB(?:\s+([\d,]+\.\d{2}))?$", line_str
      )
      if m_db:
        deb = float(m_db.group(1).replace(",", ""))
        sal = float(m_db.group(2).replace(",", "")) if m_db.group(2) else None
        tx_records.append({
            "date": active_date,
            "debet": deb,
            "kredit": 0.0,
            "saldo": sal,
        })
        continue

      # Baris Kredit (tanpa suffix DB, jika diawali tanggal)
      if m_date:
        m_cr = re.search(r"([\d,]+\.\d{2})(?:\s+([\d,]+\.\d{2}))?$", line_str)
        if m_cr:
          if m_cr.group(2):
            krd = float(m_cr.group(1).replace(",", ""))
            sal = float(m_cr.group(2).replace(",", ""))
          else:
            krd = float(m_cr.group(1).replace(",", ""))
            sal = None
          if krd > 0:
            tx_records.append({
                "date": active_date,
                "debet": 0.0,
                "kredit": krd,
                "saldo": sal,
            })

  cur_s = opening_bal
  for r in tx_records:
    if r["saldo"] is not None:
      cur_s = r["saldo"]
    elif cur_s is not None:
      cur_s = cur_s + r["kredit"] - r["debet"]
      r["saldo"] = cur_s

  for i in range(len(tx_records) - 2, -1, -1):
    if tx_records[i]["saldo"] is None and tx_records[i + 1]["saldo"] is not None:
      tx_records[i]["saldo"] = (
          tx_records[i + 1]["saldo"]
          - tx_records[i + 1]["kredit"]
          + tx_records[i + 1]["debet"]
      )

  balances = [
      r["saldo"]
      for r in tx_records
      if r["saldo"] is not None and r["saldo"] >= 1000
  ]
  if opening_bal is not None:
    balances.append(opening_bal)

  freq_db = (
      freq_db_off
      if freq_db_off is not None
      else sum(1 for r in tx_records if r["debet"] > 0)
  )
  freq_cr = (
      freq_cr_off
      if freq_cr_off is not None
      else sum(1 for r in tx_records if r["kredit"] > 0)
  )
  mutasi_db = (
      mutasi_db_off
      if mutasi_db_off is not None
      else sum(r["debet"] for r in tx_records)
  )
  mutasi_cr = (
      mutasi_cr_off
      if mutasi_cr_off is not None
      else sum(r["kredit"] for r in tx_records)
  )

  return (
      tx_records,
      balances,
      opening_bal,
      freq_db,
      freq_cr,
      mutasi_db,
      mutasi_cr,
  )


def parse_bjb_clean(pdf, all_text):
  tx_records = []
  balances = []
  opening_bal = None

  for page in pdf.pages:
    tables = page.extract_tables() or []
    for tbl in tables:
      for row in tbl:
        if not row or len(row) < 5:
          continue
        row_str = " ".join(str(c) for c in row if c)
        if (
            "Posting Date" in row_str
            or "Debit Transaction" in row_str
            or "Narasi" in row_str
        ):
          continue

        date_val = ""
        for cell in row[:4]:
          if cell:
            m_d = re.search(r"(\d{2}\s+[A-Za-z]{3})", str(cell))
            if m_d:
              date_val = m_d.group(1)
              break

        nums_in_cells = []
        for cell in row[3:]:
          if cell:
            for n in re.findall(r"[\d,]+\.\d{2}", str(cell)):
              nums_in_cells.append(float(n.replace(",", "")))

        deb_val = 0.0
        krd_val = 0.0
        sal_val = None

        if len(nums_in_cells) == 2:
          amt = nums_in_cells[0]
          sal_val = nums_in_cells[1]
          is_deb = False
          if len(row) >= 7 and row[4]:
            deb_cell = str(row[4])
            if re.search(r"[\d,]+\.\d{2}", deb_cell):
              is_deb = True
          else:
            is_deb = True
            if any(
                k in row_str.upper() for k in ["TRF DARI", "MASUK", "BUNGA"]
            ) and "PAJAK" not in row_str.upper() and "DB" not in row_str.upper():
              is_deb = False

          deb_val = amt if is_deb else 0.0
          krd_val = 0.0 if is_deb else amt

        elif len(nums_in_cells) >= 3:
          deb_val = nums_in_cells[0]
          krd_val = nums_in_cells[1]
          sal_val = nums_in_cells[2]

        if date_val and (deb_val > 0 or krd_val > 0):
          tx_records.append({
              "date": date_val,
              "debet": deb_val,
              "kredit": krd_val,
              "saldo": sal_val,
          })
          if sal_val is not None:
            balances.append(sal_val)

  # Fallback dengan regex baris teks jika tabel tidak terbaca
  if not tx_records:
    for page in pdf.pages:
      txt = page.extract_text() or ""
      for line in txt.split("\n"):
        line_str = line.strip()
        m = re.match(
            r"^(\d+)\s+(\d{2}\s+[A-Za-z]{3}\s+\d{4})\s+(\d{2}\s+[A-Za-z]{3}\s+\d{4})\s+(.+?)\s+([\d,]+\.\d{2})\s+([\d,]+\.\d{2})$",
            line_str,
        )
        if m:
          no, d1, d2, narasi, amt_s, sal_s = m.groups()
          amt = float(amt_s.replace(",", ""))
          sal = float(sal_s.replace(",", ""))
          tgl = d1[:6]
          is_deb = True
          if (
              any(k in narasi.upper() for k in ["TRF DARI", "MASUK", "BUNGA"])
              or " KR " in narasi.upper()
          ) and ("PAJAK" not in narasi.upper() and "DB" not in narasi.upper()):
            is_deb = False
          deb = amt if is_deb else 0.0
          krd = 0.0 if is_deb else amt
          tx_records.append({
              "date": tgl,
              "debet": deb,
              "kredit": krd,
              "saldo": sal,
          })
          balances.append(sal)

  if tx_records:
    first = tx_records[0]
    if first["saldo"] is not None:
      opening_bal = first["saldo"] + first["debet"] - first["kredit"]
      balances.insert(0, opening_bal)

  freq_db = sum(1 for r in tx_records if r["debet"] > 0)
  freq_cr = sum(1 for r in tx_records if r["kredit"] > 0)
  mutasi_db = sum(r["debet"] for r in tx_records)
  mutasi_cr = sum(r["kredit"] for r in tx_records)

  return (
      tx_records,
      balances,
      opening_bal,
      freq_db,
      freq_cr,
      mutasi_db,
      mutasi_cr,
  )


# ==============================================================================
# 3. DISPATCHER UTAMA
# ==============================================================================
def parse_rekening_universal(pdf_file_or_path):
  all_text = ""
  pages_text = []
  fname = (
      getattr(pdf_file_or_path, "name", "")
      if not isinstance(pdf_file_or_path, str)
      else pdf_file_or_path
  )

  with pdfplumber.open(pdf_file_or_path) as pdf:
    for p in pdf.pages:
      txt = p.extract_text() or ""
      pages_text.append(txt)
      all_text += "\n" + txt

    bank, bulan = detect_bank_and_month(all_text, fname)
    extracted_meta = extract_metadata_from_pdf(
        all_text, pages_text[0] if pages_text else "", bank
    )

    tx_records = []
    balances = []
    opening_bal = None

    if bank == "BCA":
      tx_records, balances, opening_bal, freq_db, freq_cr, mutasi_db, mutasi_cr = (
          parse_bca_clean(pdf, all_text)
      )
    elif bank == "BJB":
      tx_records, balances, opening_bal, freq_db, freq_cr, mutasi_db, mutasi_cr = (
          parse_bjb_clean(pdf, all_text)
      )
    elif bank == "BSI":
      tx_records, balances, opening_bal, freq_db, freq_cr, mutasi_db, mutasi_cr = (
          parse_bsi_clean(pdf, all_text)
      )
    elif bank == "BRI":
      for txt in pages_text:
        for line in txt.split("\n"):
          parts = line.strip().split()
          if len(parts) >= 4 and re.match(r"^\d{2}/\d{2}/\d{2}$", parts[0]):
            nums = [p.replace(",", "") for p in parts[-3:]]
            if all(re.match(r"^\d+\.\d{2}$", n) for n in nums):
              d_val, k_val, s_val = (
                  float(nums[0]),
                  float(nums[1]),
                  float(nums[2]),
              )
              tx_records.append({
                  "date": parts[0][:5],
                  "debet": d_val,
                  "kredit": k_val,
                  "saldo": s_val,
              })
              balances.append(s_val)

      tot_m = re.search(
          r"Total Transaksi Debet\s+Total Transaksi Kredit\s+Saldo Akhir\s*\n\s*([\d,\.]+)\s+([\d,\.]+)\s+([\d,\.]+)\s+([\d,\.]+)",
          all_text,
      )
      if tot_m:
        mutasi_db = float(tot_m.group(2).replace(",", ""))
        mutasi_cr = float(tot_m.group(3).replace(",", ""))
      else:
        mutasi_db = sum(r["debet"] for r in tx_records)
        mutasi_cr = sum(r["kredit"] for r in tx_records)

      freq_db = sum(1 for r in tx_records if r["debet"] > 0)
      freq_cr = sum(1 for r in tx_records if r["kredit"] > 0)
      if tx_records:
        opening_bal = (
            tx_records[0]["saldo"]
            + tx_records[0]["debet"]
            - tx_records[0]["kredit"]
        )

    elif bank == "BNI":
      text_clean = all_text.replace("|", " ")
      bni_deb = re.search(
          r"Total\s+Deb[ei]t\s*:\s*(?:(\d{1,4})\s+([\d,]+\.\d{2})|([\d,]+\.\d{2})\s+(\d{1,4}))",
          text_clean,
      )
      bni_crd = re.search(
          r"Total\s+Cr[ei]dit\s*:\s*(?:(\d{1,4})\s+([\d,]+\.\d{2})|([\d,]+\.\d{2})\s+(\d{1,4}))",
          text_clean,
      )
      tx_records, balances, opening_bal = parse_bni_clean(pdf, all_text)

      if bni_deb and bni_crd:
        freq_db = (
            int(bni_deb.group(1))
            if bni_deb.group(1)
            else int(bni_deb.group(4))
        )
        mutasi_db = (
            float(bni_deb.group(2).replace(",", ""))
            if bni_deb.group(1)
            else float(bni_deb.group(3).replace(",", ""))
        )
        freq_cr = (
            int(bni_crd.group(1))
            if bni_crd.group(1)
            else int(bni_crd.group(4))
        )
        mutasi_cr = (
            float(bni_crd.group(2).replace(",", ""))
            if bni_crd.group(1)
            else float(bni_crd.group(3).replace(",", ""))
        )
      else:
        freq_db = sum(1 for r in tx_records if r["debet"] > 0)
        freq_cr = sum(1 for r in tx_records if r["kredit"] > 0)
        mutasi_db = sum(r["debet"] for r in tx_records)
        mutasi_cr = sum(r["kredit"] for r in tx_records)

    elif bank == "PERMATA":
      tx_records, balances, opening_bal, freq_db, freq_cr, mutasi_db, mutasi_cr = (
          parse_permata_clean(pages_text, all_text)
      )
    elif bank == "MANDIRI":
      tx_records, balances, opening_bal, freq_db, freq_cr, mutasi_db, mutasi_cr = (
          parse_mandiri_clean(pages_text, all_text, pdf=pdf)
      )
    else:  # UMUM
      awal_m = re.search(r"SALDO AWAL\s*:\s*([\d,\.]+)", all_text)
      opening_bal = (
          float(awal_m.group(1).replace(",", "")) if awal_m else None
      )

      for txt in pages_text:
        for line in txt.split("\n"):
          line_clean = line.strip()
          if any(
              x in line_clean.upper()
              for x in [
                  "NO. REKENING",
                  "HALAMAN",
                  "PERIODE",
                  "MATA UANG",
                  "SALDO AWAL",
                  "MUTASI CR",
                  "MUTASI DB",
              ]
          ):
            continue

          m_db = re.search(
              r"([\d,]+\.\d{2})\s+(?:DB|DR)(?:\s+([\d,]+\.\d{2}))?", line_clean
          )
          if m_db:
            nom_db = float(m_db.group(1).replace(",", ""))
            s_val = (
                float(m_db.group(2).replace(",", "")) if m_db.group(2) else None
            )
            tgl_m = re.match(r"^(\d{2}/\d{2})", line_clean)
            tgl_str = tgl_m.group(1) if tgl_m else ""
            tx_records.append({
                "date": tgl_str,
                "debet": nom_db,
                "kredit": 0.0,
                "saldo": s_val,
            })
            continue

          amounts = re.findall(r"\b([\d,]+\.\d{2})\b", line_clean)
          if amounts:
            nums = [
                float(a.replace(",", ""))
                for a in amounts
                if float(a.replace(",", "")) < 50_000_000_000
            ]
            tgl_m = re.match(r"^(\d{2}/\d{2})", line_clean)
            tgl_str = tgl_m.group(1) if tgl_m else ""
            if len(nums) >= 2:
              nom_cr, s_val = nums[-2], nums[-1]
              tx_records.append({
                  "date": tgl_str,
                  "debet": 0.0,
                  "kredit": nom_cr,
                  "saldo": s_val,
              })
            elif len(nums) == 1 and tgl_str:
              tx_records.append({
                  "date": tgl_str,
                  "debet": 0.0,
                  "kredit": nums[0],
                  "saldo": None,
              })

      cur_s = opening_bal
      for r in tx_records:
        if r["saldo"] is not None:
          cur_s = r["saldo"]
        elif cur_s is not None:
          cur_s = cur_s + r["kredit"] - r["debet"]
          r["saldo"] = cur_s

      for i in range(len(tx_records) - 2, -1, -1):
        if (
            tx_records[i]["saldo"] is None
            and tx_records[i + 1]["saldo"] is not None
        ):
          tx_records[i]["saldo"] = (
              tx_records[i + 1]["saldo"]
              - tx_records[i + 1]["kredit"]
              + tx_records[i + 1]["debet"]
          )

      balances = [
          r["saldo"]
          for r in tx_records
          if r["saldo"] is not None and r["saldo"] >= 1000
      ]
      cr_m = re.search(r"MUTASI CR\s*:\s*([\d,\.]+)", all_text)
      db_m = re.search(r"MUTASI DB\s*:\s*([\d,\.]+)", all_text)

      freq_db = sum(1 for r in tx_records if r["debet"] > 0)
      freq_cr = sum(1 for r in tx_records if r["kredit"] > 0)
      mutasi_cr = (
          float(cr_m.group(1).replace(",", ""))
          if cr_m
          else sum(r["kredit"] for r in tx_records)
      )
      mutasi_db = (
          float(db_m.group(1).replace(",", ""))
          if db_m
          else sum(r["debet"] for r in tx_records)
      )

  valid_b = [b for b in balances if b >= 1000]
  saldo_max = max(valid_b) if valid_b else (opening_bal if opening_bal else 0.0)
  saldo_min = min(valid_b) if valid_b else (opening_bal if opening_bal else 0.0)
  saldo_avg = (
      float(np.mean(valid_b))
      if valid_b
      else (opening_bal if opening_bal else 0.0)
  )

  return {
      "bank": bank,
      "bulan": bulan,
      "opening_bal": opening_bal,
      "freq_db": freq_db,
      "freq_cr": freq_cr,
      "mutasi_db": mutasi_db,
      "mutasi_cr": mutasi_cr,
      "saldo_max": saldo_max,
      "saldo_avg": saldo_avg,
      "saldo_min": saldo_min,
      "tx_records": tx_records,
      "metadata": extracted_meta,
  }


def sort_resume_chronological(resume_list):
  filtered = [
      item
      for item in resume_list
      if item.get("bulan") and item.get("bulan") != "Bulan"
  ]

  def get_month_index(item):
    b = item.get("bulan", "")
    return MONTH_NAMES_ID.index(b) if b in MONTH_NAMES_ID else 99

  return sorted(filtered, key=get_month_index)


# ==============================================================================
# 4. GENERATOR FORM PDF & EXCEL (IN-MEMORY BYTES)
# ==============================================================================
def generate_form_pdf_bytes(
    header_info, resume_list, note_oh="", nama_so="", nama_oh=""
):
  sorted_resume = sort_resume_chronological(resume_list)
  pdf_buffer = io.BytesIO()

  doc = SimpleDocTemplate(
      pdf_buffer,
      pagesize=portrait(A4),
      rightMargin=18,
      leftMargin=18,
      topMargin=25,
      bottomMargin=25,
  )
  story = []
  styles = getSampleStyleSheet()

  title_style = ParagraphStyle(
      "DocTitle",
      parent=styles["Normal"],
      fontName="Helvetica-Bold",
      fontSize=12,
      alignment=1,
      spaceAfter=15,
  )
  story.append(
      Paragraph("<b>FORM VALIDASI MUTASI REKENING</b>", title_style)
  )

  header_rows = [
      ["Cabang", ":", header_info.get("cabang", "")],
      ["Nama Cust", ":", header_info.get("nama_cust", "")],
      ["", "", ""],
      ["Nomor Rekening", ":", header_info.get("no_rekening", "")],
      ["Nama Bank", ":", header_info.get("nama_bank", "")],
      [
          "Nama Pemegang Rekening",
          ":",
          header_info.get("nama_pemegang_rek", ""),
      ],
  ]

  t_header = Table(header_rows, colWidths=[140, 15, 404])
  t_header.setStyle(
      TableStyle([
          ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
          ("FONTSIZE", (0, 0), (-1, -1), 8.5),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
          ("TOPPADDING", (0, 0), (-1, -1), 2),
          ("BACKGROUND", (2, 0), (2, 1), colors.HexColor("#E2EFDA")),
          ("BACKGROUND", (2, 3), (2, 5), colors.HexColor("#E2EFDA")),
          ("BOX", (2, 0), (2, 1), 0.5, colors.black),
          ("BOX", (2, 3), (2, 5), 0.5, colors.black),
      ])
  )
  story.append(t_header)
  story.append(Spacer(1, 12))

  sec_title_style = ParagraphStyle(
      "SecTitle",
      parent=styles["Normal"],
      fontName="Helvetica-Bold",
      fontSize=9.5,
      spaceAfter=4,
  )
  story.append(Paragraph("<b>Resume Mutasi Rekening</b>", sec_title_style))

  table_data = [
      ["Bulan", "Frekuensi", "", "Mutasi", "", "Saldo", "", ""],
      [
          "",
          "Debet",
          "Kredit",
          "Debet",
          "Kredit",
          "Tertinggi",
          "Rata-Rata",
          "Terendah",
      ],
  ]

  totals = {
      "f_db": 0,
      "f_cr": 0,
      "m_db": 0.0,
      "m_cr": 0.0,
      "s_max": 0.0,
      "s_avg": 0.0,
      "s_min": 0.0,
  }
  n = len(sorted_resume)

  for item in sorted_resume:
    totals["f_db"] += item.get("freq_db", 0)
    totals["f_cr"] += item.get("freq_cr", 0)
    totals["m_db"] += item.get("mutasi_db", 0.0)
    totals["m_cr"] += item.get("mutasi_cr", 0.0)
    totals["s_max"] += item.get("saldo_max", 0.0)
    totals["s_avg"] += item.get("saldo_avg", 0.0)
    totals["s_min"] += item.get("saldo_min", 0.0)

    table_data.append([
        item.get("bulan", "-"),
        f"{item.get('freq_db', 0):,}",
        f"{item.get('freq_cr', 0):,}",
        f"{item.get('mutasi_db', 0.0):,.2f}",
        f"{item.get('mutasi_cr', 0.0):,.2f}",
        f"{item.get('saldo_max', 0.0):,.2f}",
        f"{item.get('saldo_avg', 0.0):,.2f}",
        f"{item.get('saldo_min', 0.0):,.2f}",
    ])

  table_data.append([
      "Rata-Rata",
      f"{int(totals['f_db']/n):,}" if n else "0",
      f"{int(totals['f_cr']/n):,}" if n else "0",
      f"{totals['m_db']/n:,.2f}" if n else "0.00",
      f"{totals['m_cr']/n:,.2f}" if n else "0.00",
      f"{totals['s_max']/n:,.2f}" if n else "0.00",
      f"{totals['s_avg']/n:,.2f}" if n else "0.00",
      f"{totals['s_min']/n:,.2f}" if n else "0.00",
  ])

  col_widths = [56, 38, 38, 86, 86, 85, 85, 85]
  t_resume = Table(table_data, colWidths=col_widths)
  t_resume.setStyle(
      TableStyle([
          ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#D9E1F2")),
          ("ALIGN", (0, 0), (-1, -1), "CENTER"),
          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("FONTNAME", (0, 0), (-1, 1), "Helvetica-Bold"),
          ("FONTSIZE", (0, 0), (-1, -1), 7.2),
          ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
          ("SPAN", (0, 0), (0, 1)),
          ("SPAN", (1, 0), (2, 0)),
          ("SPAN", (3, 0), (4, 0)),
          ("SPAN", (5, 0), (7, 0)),
          ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2EFDA")),
          ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
          ("ALIGN", (3, 2), (-1, -1), "RIGHT"),
          ("RIGHTPADDING", (3, 2), (-1, -1), 3),
          ("LEFTPADDING", (3, 2), (-1, -1), 3),
      ])
  )
  story.append(t_resume)
  story.append(Spacer(1, 15))

  story.append(Paragraph("<b>Note OH:</b>", sec_title_style))
  note_box = Table([[note_oh if note_oh else "-"]], colWidths=[559])
  note_box.setStyle(
      TableStyle([
          ("BOX", (0, 0), (-1, -1), 0.5, colors.black),
          ("MINROWHEIGHT", (0, 0), (-1, -1), 35),
          ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
          ("FONTSIZE", (0, 0), (-1, -1), 8),
          ("VALIGN", (0, 0), (-1, -1), "TOP"),
      ])
  )
  story.append(note_box)
  story.append(Spacer(1, 30))

  so_display = f"SO: {nama_so}" if nama_so else "SO:"
  oh_display = f"Operation Head: {nama_oh}" if nama_oh else "Operation Head:"
  sign_data = [
      ["Mengajukan,", "Menyetujui,"],
      ["", ""],
      ["", ""],
      ["", ""],
      [so_display, oh_display],
  ]
  t_sign = Table(sign_data, colWidths=[275, 284])
  t_sign.setStyle(
      TableStyle([
          ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
          ("FONTSIZE", (0, 0), (-1, -1), 9),
          ("VALIGN", (0, 0), (-1, -1), "TOP"),
      ])
  )
  story.append(t_sign)

  doc.build(story)
  pdf_buffer.seek(0)
  return pdf_buffer


def generate_form_excel_bytes(
    header_info, resume_list, note_oh="", nama_so="", nama_oh=""
):
  sorted_resume = sort_resume_chronological(resume_list)

  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Validasi Mutasi"
  ws.views.sheetView[0].showGridLines = True

  font_title = Font(name="Arial", size=11, bold=True)
  font_sec = Font(name="Arial", size=10, bold=True)
  font_bold = Font(name="Arial", size=9, bold=True)
  font_norm = Font(name="Arial", size=9)

  fill_green = PatternFill(
      start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
  )
  fill_blue = PatternFill(
      start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
  )
  fill_yellow = PatternFill(
      start_color="FFFF00", end_color="FFFF00", fill_type="solid"
  )
  fill_grey = PatternFill(
      start_color="808080", end_color="808080", fill_type="solid"
  )
  fill_soft_blue = PatternFill(
      start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"
  )

  thin = Side(border_style="thin", color="000000")
  box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

  ws.merge_cells("C1:G1")
  ws["C1"] = "FORM VALIDASI MUTASI REKENING"
  ws["C1"].font = font_title
  ws["C1"].alignment = Alignment(horizontal="center", vertical="center")

  headers = [
      (3, "Cabang", header_info.get("cabang", "")),
      (4, "Nama Cust", header_info.get("nama_cust", "")),
      (6, "Nomor Rekening", header_info.get("no_rekening", "")),
      (7, "Nama Bank", header_info.get("nama_bank", "")),
      (8, "Nama Pemegang Rekening", header_info.get("nama_pemegang_rek", "")),
  ]

  for r, label, val in headers:
    ws.cell(row=r, column=1, value=label).font = font_norm
    ws.cell(row=r, column=3, value=":").font = font_norm
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=7)
    cell = ws.cell(row=r, column=4, value=val)
    cell.font = font_norm
    cell.fill = fill_green
    for c in range(4, 8):
      ws.cell(row=r, column=c).border = box_border

  ws["A10"] = "Resume Mutasi Rekening"
  ws["A10"].font = font_sec

  ws.merge_cells("A11:A12")
  ws["A11"] = "Bulan"
  ws.merge_cells("B11:C11")
  ws["B11"] = "Frekuensi"
  ws["B12"] = "Debet"
  ws["C12"] = "Kredit"
  ws.merge_cells("D11:E11")
  ws["D11"] = "Mutasi"
  ws["D12"] = "Debet"
  ws["E12"] = "Kredit"
  ws.merge_cells("F11:H11")
  ws["F11"] = "Saldo"
  ws["F12"] = "Tertinggi"
  ws["G12"] = "Rata-Rata"
  ws["H12"] = "Terendah"

  for r in range(11, 13):
    for c in range(1, 9):
      cell = ws.cell(row=r, column=c)
      cell.fill = fill_blue
      cell.font = font_bold
      cell.alignment = Alignment(horizontal="center", vertical="center")
      cell.border = box_border

  start_r = 13
  n = len(sorted_resume)
  sum_vals = {
      "f_db": 0,
      "f_cr": 0,
      "m_db": 0.0,
      "m_cr": 0.0,
      "s_max": 0.0,
      "s_avg": 0.0,
      "s_min": 0.0,
  }

  for i, item in enumerate(sorted_resume):
    curr_r = start_r + i
    b_name = item.get("bulan", "-")

    c1 = ws.cell(row=curr_r, column=1, value=b_name)
    c1.alignment = Alignment(horizontal="center")
    c1.font = font_norm
    c1.border = box_border

    f_db = item.get("freq_db", 0)
    f_cr = item.get("freq_cr", 0)
    m_db = item.get("mutasi_db", 0.0)
    m_cr = item.get("mutasi_cr", 0.0)
    s_max = item.get("saldo_max", 0.0)
    s_avg = item.get("saldo_avg", 0.0)
    s_min = item.get("saldo_min", 0.0)

    sum_vals["f_db"] += f_db
    sum_vals["f_cr"] += f_cr
    sum_vals["m_db"] += m_db
    sum_vals["m_cr"] += m_cr
    sum_vals["s_max"] += s_max
    sum_vals["s_avg"] += s_avg
    sum_vals["s_min"] += s_min

    ws.cell(row=curr_r, column=2, value=f_db).number_format = "#,##0"
    ws.cell(row=curr_r, column=3, value=f_cr).number_format = "#,##0"
    ws.cell(row=curr_r, column=4, value=m_db).number_format = "#,##0.00"
    ws.cell(row=curr_r, column=5, value=m_cr).number_format = "#,##0.00"
    ws.cell(row=curr_r, column=6, value=s_max).number_format = "#,##0.00"
    ws.cell(row=curr_r, column=7, value=s_avg).number_format = "#,##0.00"
    ws.cell(row=curr_r, column=8, value=s_min).number_format = "#,##0.00"

    for c_idx in range(2, 9):
      cell = ws.cell(row=curr_r, column=c_idx)
      cell.font = font_norm
      cell.border = box_border
      cell.alignment = Alignment(horizontal="right")

  end_data_r = start_r + len(sorted_resume) - 1
  avg_r = end_data_r + 1

  ws.cell(row=avg_r, column=1, value="Rata-Rata")

  avg_f_db = round(sum_vals["f_db"] / n) if n else 0
  avg_f_cr = round(sum_vals["f_cr"] / n) if n else 0
  avg_m_db = (sum_vals["m_db"] / n) if n else 0.0
  avg_m_cr = (sum_vals["m_cr"] / n) if n else 0.0
  avg_s_max = (sum_vals["s_max"] / n) if n else 0.0
  avg_s_avg = (sum_vals["s_avg"] / n) if n else 0.0
  avg_s_min = (sum_vals["s_min"] / n) if n else 0.0

  ws.cell(row=avg_r, column=2, value=avg_f_db).number_format = "#,##0"
  ws.cell(row=avg_r, column=3, value=avg_f_cr).number_format = "#,##0"
  ws.cell(row=avg_r, column=4, value=avg_m_db).number_format = "#,##0.00"
  ws.cell(row=avg_r, column=5, value=avg_m_cr).number_format = "#,##0.00"
  ws.cell(row=avg_r, column=6, value=avg_s_max).number_format = "#,##0.00"
  ws.cell(row=avg_r, column=7, value=avg_s_avg).number_format = "#,##0.00"
  ws.cell(row=avg_r, column=8, value=avg_s_min).number_format = "#,##0.00"

  for c in range(1, 9):
    cell = ws.cell(row=avg_r, column=c)
    cell.fill = fill_green
    cell.font = font_bold
    cell.border = box_border
    if c == 1:
      cell.alignment = Alignment(horizontal="center")
    else:
      cell.alignment = Alignment(horizontal="right")

  note_label_r = avg_r + 2
  ws.cell(row=note_label_r, column=1, value="Note OH:").font = font_sec
  ws.merge_cells(
      start_row=note_label_r + 1,
      start_column=1,
      end_row=note_label_r + 2,
      end_column=8,
  )
  note_c = ws.cell(
      row=note_label_r + 1, column=1, value=note_oh if note_oh else "-"
  )
  note_c.font = font_norm
  note_c.alignment = Alignment(vertical="top")
  for r in range(note_label_r + 1, note_label_r + 3):
    for c in range(1, 9):
      ws.cell(row=r, column=c).border = box_border

  sign_r = note_label_r + 4
  ws.cell(row=sign_r, column=1, value="Mengajukan,").font = font_norm
  ws.cell(row=sign_r, column=6, value="Menyetujui,").font = font_norm
  ws.cell(row=sign_r + 4, column=1, value=f"SO: {nama_so}" if nama_so else "SO:").font = font_norm
  ws.cell(row=sign_r + 4, column=6, value=f"Operation Head: {nama_oh}" if nama_oh else "Operation Head:").font = font_norm

  # Kolom rincian transaksi samping
  start_col = 10
  for idx, item in enumerate(sorted_resume):
    b_name = item.get("bulan", f"Bulan {idx+1}")
    col_b = start_col + (idx * 5)

    ws.cell(row=1, column=col_b, value="Bulan")
    ws.merge_cells(
        start_row=1, start_column=col_b + 1, end_row=1, end_column=col_b + 2
    )
    ws.cell(row=1, column=col_b + 1, value=f"Mutasi {b_name}")
    ws.cell(row=1, column=col_b + 3, value="Saldo")

    ws.cell(row=2, column=col_b + 1, value="Debet")
    ws.cell(row=2, column=col_b + 2, value="Kredit")

    for r in range(1, 3):
      for c in range(col_b, col_b + 4):
        cell = ws.cell(row=r, column=c)
        cell.font = font_bold
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = box_border

    ws.cell(row=3, column=col_b, value=b_name).fill = fill_yellow
    ws.cell(row=3, column=col_b).font = font_bold
    ws.cell(row=3, column=col_b).border = box_border

    col_deb = get_column_letter(col_b + 1)
    col_krd = get_column_letter(col_b + 2)

    records = item.get("tx_records", [])
    max_tx_row = max(len(records) + 3, 40)

    ws.cell(
        row=3, column=col_b + 1, value=f"=SUM({col_deb}4:{col_deb}{max_tx_row})"
    )
    ws.cell(
        row=3, column=col_b + 2, value=f"=SUM({col_krd}4:{col_krd}{max_tx_row})"
    )

    op_bal = item.get("opening_bal")
    if op_bal is None and records:
      op_bal = (
          records[0]["saldo"] + records[0]["debet"] - records[0]["kredit"]
      )

    ws.cell(row=3, column=col_b + 3, value=op_bal)

    for c in range(col_b, col_b + 4):
      cell = ws.cell(row=3, column=c)
      cell.fill = fill_yellow
      cell.font = font_bold
      cell.border = box_border
      if c > col_b:
        cell.alignment = Alignment(horizontal="right")
        cell.number_format = "#,##0.00"

    for r_idx in range(4, max_tx_row + 1):
      tx_idx = r_idx - 4
      rec = records[tx_idx] if tx_idx < len(records) else None

      c_bln = ws.cell(
          row=r_idx, column=col_b, value=rec["date"] if rec else ""
      )
      c_bln.fill = fill_grey
      c_bln.font = Font(name="Arial", size=8, color="FFFFFF")
      c_bln.border = box_border
      c_bln.alignment = Alignment(horizontal="center")

      c_deb = ws.cell(
          row=r_idx,
          column=col_b + 1,
          value=rec["debet"] if (rec and rec["debet"] > 0) else None,
      )
      c_deb.fill = fill_yellow
      c_deb.font = font_norm
      c_deb.border = box_border
      c_deb.number_format = "#,##0.00"

      c_krd = ws.cell(
          row=r_idx,
          column=col_b + 2,
          value=rec["kredit"] if (rec and rec["kredit"] > 0) else None,
      )
      c_krd.fill = fill_yellow
      c_krd.font = font_norm
      c_krd.border = box_border
      c_krd.number_format = "#,##0.00"

      c_sal = ws.cell(
          row=r_idx, column=col_b + 3, value=rec["saldo"] if rec else None
      )
      c_sal.fill = fill_soft_blue
      c_sal.font = font_norm
      c_sal.border = box_border
      c_sal.number_format = "#,##0.00"
      c_sal.alignment = Alignment(horizontal="right")

    col_sep = col_b + 4
    for r_sep in range(1, max_tx_row + 1):
      ws.cell(row=r_sep, column=col_sep).fill = PatternFill(
          start_color="000000", end_color="000000", fill_type="solid"
      )
    ws.column_dimensions[get_column_letter(col_sep)].width = 2.5
    ws.column_dimensions[get_column_letter(col_b)].width = 12
    ws.column_dimensions[get_column_letter(col_b + 1)].width = 16
    ws.column_dimensions[get_column_letter(col_b + 2)].width = 16
    ws.column_dimensions[get_column_letter(col_b + 3)].width = 18

  for col in ["A", "B", "C", "D", "E", "F", "G", "H"]:
    ws.column_dimensions[col].width = 15
  ws.column_dimensions["A"].width = 14
  ws.column_dimensions["B"].width = 10
  ws.column_dimensions["C"].width = 10

  xlsx_buffer = io.BytesIO()
  wb.save(xlsx_buffer)
  xlsx_buffer.seek(0)
  return xlsx_buffer


# ==============================================================================
# ==============================================================================
# 5. ANTARMUKA PENGGUNA STREAMLIT (MODERN & MINIMALIS)
# ==============================================================================
st.set_page_config(
    page_title="Validasi Mutasi Rekening",
    page_icon="💳",
    layout="wide",
    initial_sidebar_state="expanded",
)

# Custom CSS Modern & Minimalis
st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700&display=swap');

    html, body, [class*="css"] {
        font-family: 'Plus Jakarta Sans', sans-serif;
    }

    /* Minimalist header adjustments */
    .block-container {
        padding-top: 1.8rem;
        padding-bottom: 3rem;
        max-width: 1200px;
    }

    /* Hero Card */
    .hero-card {
        background: linear-gradient(135deg, #ffffff 0%, #f8fafc 100%);
        border: 1px solid #e2e8f0;
        border-radius: 16px;
        padding: 24px 28px;
        margin-bottom: 24px;
        box-shadow: 0 1px 3px rgba(0,0,0,0.02), 0 6px 16px rgba(0,0,0,0.02);
    }
    .hero-badge {
        display: inline-block;
        padding: 4px 10px;
        font-size: 11px;
        font-weight: 700;
        letter-spacing: 0.05em;
        text-transform: uppercase;
        color: #2563eb;
        background: #eff6ff;
        border-radius: 9999px;
        margin-bottom: 10px;
    }
    .hero-title {
        font-size: 26px;
        font-weight: 700;
        color: #0f172a;
        margin: 0 0 6px 0;
        letter-spacing: -0.02em;
    }
    .hero-subtitle {
        font-size: 14px;
        color: #64748b;
        margin: 0 0 16px 0;
        line-height: 1.5;
    }
    .bank-tags {
        display: flex;
        flex-wrap: wrap;
        gap: 6px;
    }
    .bank-pill {
        background: #f1f5f9;
        color: #334155;
        font-size: 12px;
        font-weight: 600;
        padding: 3px 10px;
        border-radius: 6px;
        border: 1px solid #e2e8f0;
    }

    /* Sidebar minimalis */
    [data-testid="stSidebar"] {
        background-color: #f8fafc;
        border-right: 1px solid #e2e8f0;
    }
    .sidebar-section-title {
        font-size: 11px;
        font-weight: 700;
        text-transform: uppercase;
        letter-spacing: 0.08em;
        color: #64748b;
        margin: 16px 0 8px 0;
        padding-bottom: 4px;
        border-bottom: 1px solid #e2e8f0;
    }

    /* Metric Cards */
    .metric-container {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(210px, 1fr));
        gap: 14px;
        margin-bottom: 24px;
    }
    .metric-card {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 12px;
        padding: 16px 18px;
        box-shadow: 0 1px 2px rgba(0,0,0,0.03);
    }
    .metric-label {
        font-size: 12px;
        font-weight: 600;
        color: #64748b;
        margin-bottom: 6px;
    }
    .metric-value {
        font-size: 20px;
        font-weight: 700;
        color: #0f172a;
    }
    .metric-badge-debet {
        color: #dc2626;
    }
    .metric-badge-kredit {
        color: #16a34a;
    }

    /* Download Cards */
    .dl-card {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 12px;
        padding: 18px;
        height: 100%;
        box-shadow: 0 1px 2px rgba(0,0,0,0.03);
        margin-bottom: 12px;
    }
    .dl-card-title {
        font-size: 14px;
        font-weight: 600;
        color: #0f172a;
        margin-bottom: 4px;
    }
    .dl-card-desc {
        font-size: 12px;
        color: #64748b;
        margin-bottom: 14px;
    }

    /* Clean Buttons */
    div.stButton > button:first-child {
        border-radius: 8px;
        font-weight: 600;
        transition: all 0.15s ease-in-out;
    }
    div.stDownloadButton > button:first-child {
        border-radius: 8px;
        font-weight: 600;
        border: 1px solid #cbd5e1;
        transition: all 0.15s ease-in-out;
    }
    div.stDownloadButton > button:first-child:hover {
        border-color: #2563eb;
        color: #2563eb;
        background-color: #f8fafc;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# Hero / Header Modern
st.markdown(
    """
    <div class="hero-card">
        <div class="hero-badge">Financial Reconciliation System</div>
        <div class="hero-title">Form Validasi Mutasi Rekening</div>
        <div class="hero-subtitle">Ekstraksi berkas rekening koran multi-bank otomatis dan terstruktur dengan kalkulasi presisi untuk validasi dokumen finansial.</div>
        <div class="bank-tags">
            <span class="bank-pill">Mandiri</span>
            <span class="bank-pill">BCA</span>
            <span class="bank-pill">BNI</span>
            <span class="bank-pill">BRI</span>
            <span class="bank-pill">BSI</span>
            <span class="bank-pill">BJB</span>
            <span class="bank-pill">Permata</span>
            <span class="bank-pill">Nobu</span>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

# Sidebar: Form Informasi
with st.sidebar:
  st.markdown("### ⚙️ Parameter Dokumen")

  st.markdown(
      '<div class="sidebar-section-title">Informasi Nasabah</div>',
      unsafe_allow_html=True,
  )
  cabang = st.text_input("Cabang", value="", placeholder="Contoh: Jakarta Pusat")
  nama_cust = st.text_input("Nama Cust", value="", placeholder="Nama lengkap customer")
  no_rek = st.text_input("Nomor Rekening", value="", placeholder="Contoh: 1230009876543")
  nama_bank = st.selectbox(
      "Nama Bank", ["Mandiri", "BCA", "BNI", "BRI", "BSI", "BJB", "Permata", "Nobu"]
  )
  pemegang_rek = st.text_input(
      "Nama Pemegang Rekening", value="", placeholder="Sesuai buku tabungan"
  )

  st.markdown(
      '<div class="sidebar-section-title">Pengesahan & Catatan</div>',
      unsafe_allow_html=True,
  )
  nama_so = st.text_input(
      "Nama SO (Sales Officer)", value="", placeholder="Nama Sales Officer"
  )
  nama_oh = st.text_input(
      "Operation Head", value="", placeholder="Nama Operation Head"
  )
  note_oh = st.text_area(
      "Note OH", value="-", placeholder="Catatan dari Operation Head..."
  )

# Upload Container
st.markdown("#### 📂 Berkas Rekening Koran")
uploaded_files = st.file_uploader(
    "Pilih 1 hingga 3 Berkas Rekening Koran (Format PDF)",
    type=["pdf"],
    accept_multiple_files=True,
    help="Unggah dokumen PDF mutasi rekening untuk dianalisis dan direkapitulasi secara otomatis.",
)

if uploaded_files:
  col_btn, _ = st.columns([1, 2])
  with col_btn:
    proses_clicked = st.button("🚀 Ekstraksi & Rekap Data", type="primary", use_container_width=True)

  if proses_clicked:
    resume_list = []
    progress_bar = st.progress(0, text="Memulai ekstraksi...")

    for i, file_obj in enumerate(uploaded_files):
      progress_bar.progress(
          (i + 1) / len(uploaded_files),
          text=f"Membaca berkas: {file_obj.name}...",
      )
      extracted = parse_rekening_universal(file_obj)
      resume_list.append(extracted)

    st.session_state["resume_list"] = resume_list
    progress_bar.empty()
    st.toast("Ekstraksi data mutasi berhasil selesai!", icon="✅")

# Pratinjau & Unduh Dokumen
if "resume_list" in st.session_state and st.session_state["resume_list"]:
  res_list = st.session_state["resume_list"]
  sorted_data = sort_resume_chronological(res_list)

  # Ambil metadata dari file hasil ekstraksi jika input di sidebar masih kosong
  detected_meta = {}
  for item in sorted_data:
    if item.get("metadata"):
      detected_meta = item["metadata"]
      break

  final_cabang = cabang.strip() if cabang.strip() else detected_meta.get("cabang", "")
  final_cust = nama_cust.strip() if nama_cust.strip() else detected_meta.get("nama_cust", "")
  final_rek = no_rek.strip() if no_rek.strip() else detected_meta.get("no_rekening", "")
  final_bank = nama_bank if (nama_bank and nama_bank != "Mandiri") else (detected_meta.get("nama_bank") or nama_bank)
  final_pemegang = pemegang_rek.strip() if pemegang_rek.strip() else (final_cust or detected_meta.get("nama_pemegang_rek", ""))

  header_input = {
      "cabang": final_cabang,
      "nama_cust": final_cust,
      "no_rekening": final_rek,
      "nama_bank": final_bank,
      "nama_pemegang_rek": final_pemegang,
      "nama_so": nama_so,
      "nama_oh": nama_oh,
  }

  if detected_meta.get("nama_cust") or detected_meta.get("no_rekening"):
    st.info(
        f"📋 **Data Rekening Terdeteksi Otomatis:** Bank: **{final_bank}** | No. Rek: **{final_rek}** | Nasabah: **{final_cust}** | Cabang: **{final_cabang}**"
    )

  # Ringkasan KPI
  total_debet = sum(r.get("mutasi_db", 0) for r in sorted_data)
  total_kredit = sum(r.get("mutasi_cr", 0) for r in sorted_data)
  avg_saldos = [
      r.get("saldo_avg", 0) for r in sorted_data if r.get("saldo_avg")
  ]
  overall_avg = sum(avg_saldos) / len(avg_saldos) if avg_saldos else 0

  st.markdown("---")
  st.markdown("#### 📊 Ringkasan Mutasi")

  m1, m2, m3, m4 = st.columns(4)
  with m1:
    st.metric("Periode Diproses", f"{len(sorted_data)} Bulan")
  with m2:
    st.metric("Total Debet", f"Rp {total_debet:,.2f}")
  with m3:
    st.metric("Total Kredit", f"Rp {total_kredit:,.2f}")
  with m4:
    st.metric("Rata-Rata Saldo", f"Rp {overall_avg:,.2f}")

  # Tabel Pratinjau
  preview_rows = []
  for r in sorted_data:
    preview_rows.append({
        "Bulan": r.get("bulan"),
        "Bank": r.get("bank"),
        "Freq Debet": r.get("freq_db"),
        "Freq Kredit": r.get("freq_cr"),
        "Total Debet (Rp)": f"{r.get('mutasi_db', 0):,.2f}",
        "Total Kredit (Rp)": f"{r.get('mutasi_cr', 0):,.2f}",
        "Saldo Tertinggi (Rp)": f"{r.get('saldo_max', 0):,.2f}",
        "Saldo Rata-Rata (Rp)": f"{r.get('saldo_avg', 0):,.2f}",
        "Saldo Terendah (Rp)": f"{r.get('saldo_min', 0):,.2f}",
    })
  st.dataframe(preview_rows, use_container_width=True, hide_index=True)

  # Generate Buffer File
  pdf_bytes = generate_form_pdf_bytes(
      header_input,
      res_list,
      note_oh=note_oh,
      nama_so=nama_so,
      nama_oh=nama_oh,
  )
  xlsx_bytes = generate_form_excel_bytes(
      header_input,
      res_list,
      note_oh=note_oh,
      nama_so=nama_so,
      nama_oh=nama_oh,
  )

  # Unduh Hasil
  st.markdown("#### 📥 Unduh Dokumen Form Validasi")
  col_dl1, col_dl2 = st.columns(2)

  with col_dl1:
    st.markdown(
        """
        <div class="dl-card">
            <div class="dl-card-title">📄 Dokumen PDF Resmi</div>
            <div class="dl-card-desc">Format standar A4 siap cetak dengan tanda tangan SO & Operation Head.</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.download_button(
        label="Unduh PDF Form Validasi",
        data=pdf_bytes,
        file_name="FORM_VALIDASI_MUTASI_REKENING.pdf",
        mime="application/pdf",
        use_container_width=True,
    )

  with col_dl2:
    st.markdown(
        """
        <div class="dl-card">
            <div class="dl-card-title">📊 Dokumen Excel Spreadsheet</div>
            <div class="dl-card-desc">Format lembar kerja dinamis lengkap dengan rincian transaksi per bulan.</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.download_button(
        label="Unduh Excel Form Validasi (.xlsx)",
        data=xlsx_bytes,
        file_name="FORM_VALIDASI_MUTASI_REKENING.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )